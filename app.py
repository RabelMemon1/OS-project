from flask import Flask, jsonify, render_template, send_file
import subprocess
import os

app = Flask(__name__)

# ===================================================
# HELPER FUNCTION - Bash script run karne ke liye
# ===================================================


def run_script(script_name, args=None):
    script_path = os.path.join(os.path.dirname(
        __file__), 'scripts', script_name)
    cmd = ['bash', script_path]
    if args:
        cmd.extend(args)
    result = subprocess.run(cmd, capture_output=True, text=True)
    return result.stdout.strip()

# ===================================================
# MAIN ROUTE - Dashboard
# ===================================================


@app.route('/')
def index():
    return render_template('index.html')

# ===================================================
# API ROUTES
# ===================================================

# CPU Data


@app.route('/api/cpu')
def cpu():
    try:
        output = run_script('cpu.sh')
        import json
        data = json.loads(output)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Memory Data


@app.route('/api/memory')
def memory():
    try:
        output = run_script('memory.sh')
        import json
        data = json.loads(output)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Disk Data


@app.route('/api/disk')
def disk():
    try:
        output = run_script('disk.sh')
        import json
        data = json.loads(output)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Process Data


@app.route('/api/processes')
def processes():
    try:
        output = run_script('process.sh')
        import json
        data = json.loads(output)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Kill Process


@app.route('/api/kill/<int:pid>', methods=['POST'])
def kill_process(pid):
    try:
        output = run_script('kill_process.sh', [str(pid)])
        import json
        data = json.loads(output)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ===================================================
# Generate Report - Word Document (.docx)
# Redesigned: KPI summary cards, theme-colored tables,
# full-row highlighting for WARNING/CRITICAL disk rows,
# zebra-striped key/value + process tables.
# ===================================================


@app.route('/api/report')
def report():
    try:
        import json
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from datetime import datetime
        import socket

        # Collect data from bash scripts
        cpu_data = json.loads(run_script('cpu.sh'))
        mem_data = json.loads(run_script('memory.sh'))
        disk_data = json.loads(run_script('disk.sh'))
        proc_data = json.loads(run_script('process.sh'))

        # System info
        hostname = socket.gethostname()
        os_raw = subprocess.run(
            ['cat', '/etc/os-release'], capture_output=True, text=True).stdout
        os_name = next((l.split('=')[1].strip().strip('"') for l in os_raw.split(
            '\n') if l.startswith('PRETTY_NAME')), 'Linux')
        kernel = subprocess.run(
            ['uname', '-r'], capture_output=True, text=True).stdout.strip()
        uptime = subprocess.run(
            ['uptime', '-p'], capture_output=True, text=True).stdout.strip()
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        # ── Theme ────────────────────────────────────
        NAVY = '1F3864'
        BLUE = '2E75B6'
        GREEN = '1E8449'
        GREEN_BG = 'E8F5E9'
        AMBER = 'B7791F'
        AMBER_BG = 'FFF8E1'
        RED = 'C0392B'
        RED_BG = 'FDEDEC'
        GRAY = '757575'
        ROW_ALT = 'F5F7FA'

        def status_colors(status):
            if status == 'CRITICAL':
                return RED, RED_BG
            if status == 'WARNING':
                return AMBER, AMBER_BG
            return GREEN, GREEN_BG

        def hex_rgb(h):
            return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

        doc = Document()

        # Base font for the whole document
        base_style = doc.styles['Normal']
        base_style.font.name = 'Calibri'
        base_style.font.size = Pt(11)

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # ── Helpers ──────────────────────────────────
        def set_cell_bg(cell, hex_color):
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), hex_color)
            tcPr.append(shd)

        def set_table_borders(table, color='D9D9D9', sz=4):
            tblPr = table._tbl.tblPr
            borders = OxmlElement('w:tblBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                el = OxmlElement(f'w:{edge}')
                el.set(qn('w:val'), 'single')
                el.set(qn('w:sz'), str(sz))
                el.set(qn('w:space'), '0')
                el.set(qn('w:color'), color)
                borders.append(el)
            tblPr.append(borders)

        def set_col_widths(table, widths_in):
            table.autofit = False
            for row in table.rows:
                for cell, w in zip(row.cells, widths_in):
                    cell.width = Inches(w)

        def vcenter(cell):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        def cell_text(cell, text, size=10, bold=False, color='000000', align=None):
            cell.text = ''
            p = cell.paragraphs[0]
            if align:
                p.alignment = align
            run = p.add_run(text)
            run.font.size = Pt(size)
            run.bold = bold
            run.font.color.rgb = hex_rgb(color)
            vcenter(cell)
            return run

        def add_heading(text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = hex_rgb(NAVY)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bot = OxmlElement('w:bottom')
            bot.set(qn('w:val'), 'single')
            bot.set(qn('w:sz'), '8')
            bot.set(qn('w:space'), '2')
            bot.set(qn('w:color'), BLUE)
            pBdr.append(bot)
            pPr.append(pBdr)

        def add_kv_table(pairs):
            """Two-column, zebra-striped key/value block.
            A 'Status' row is auto-colored per status value."""
            table = doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for i, (k, v) in enumerate(pairs):
                row = table.add_row()
                bg = 'FFFFFF' if i % 2 == 0 else ROW_ALT
                c0, c1 = row.cells
                cell_text(c0, k, size=10.5, bold=True, color='555555')
                if k == 'Status':
                    fg, _ = status_colors(str(v))
                    cell_text(c1, str(v), size=10.5, bold=True, color=fg)
                else:
                    cell_text(c1, str(v), size=10.5,
                              bold=False, color='1A1A1A')
                set_cell_bg(c0, bg)
                set_cell_bg(c1, bg)
            set_col_widths(table, [1.8, 4.2])
            return table

        def add_kpi_row(items):
            """items: list of (label, value_str, status) -> colored KPI cards."""
            table = doc.add_table(rows=1, cols=len(items))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_col_widths(table, [1.8] * len(items))
            for cell, (label, value, status) in zip(table.rows[0].cells, items):
                fg, bg = status_colors(status)
                set_cell_bg(cell, bg)
                cell.text = ''
                p1 = cell.paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p1.paragraph_format.space_before = Pt(6)
                r1 = p1.add_run(value)
                r1.bold = True
                r1.font.size = Pt(22)
                r1.font.color.rgb = hex_rgb(fg)

                p2 = cell.add_paragraph()
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r2 = p2.add_run(label)
                r2.font.size = Pt(9)
                r2.font.color.rgb = hex_rgb('555555')

                p3 = cell.add_paragraph()
                p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p3.paragraph_format.space_after = Pt(6)
                r3 = p3.add_run(status)
                r3.bold = True
                r3.font.size = Pt(9)
                r3.font.color.rgb = hex_rgb(fg)
                vcenter(cell)
            set_table_borders(table, color='D9D9D9', sz=4)
            return table

        def make_header_row(table, headers, widths=None):
            row = table.rows[0]
            for i, h in enumerate(headers):
                cell = row.cells[i]
                cell_text(cell, h, size=9.5, bold=True, color='FFFFFF',
                          align=WD_ALIGN_PARAGRAPH.CENTER)
                set_cell_bg(cell, NAVY)
            if widths:
                set_col_widths(table, widths)

        # ── Title ────────────────────────────────────
        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tr = t.add_run('LINUX SYSTEM MONITORING TOOL')
        tr.bold = True
        tr.font.size = Pt(21)
        tr.font.color.rgb = hex_rgb(NAVY)

        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = s.add_run('System Performance Report')
        sr.font.size = Pt(13)
        sr.font.color.rgb = hex_rgb(BLUE)

        # separator line
        sep = doc.add_paragraph()
        sep.paragraph_format.space_before = Pt(6)
        sep.paragraph_format.space_after = Pt(2)
        pPr = sep._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single')
        bot.set(qn('w:sz'), '18')
        bot.set(qn('w:space'), '1')
        bot.set(qn('w:color'), NAVY)
        pBdr.append(bot)
        pPr.append(pBdr)

        d = doc.add_paragraph()
        d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        d.paragraph_format.space_before = Pt(6)
        dr = d.add_run(f'Generated: {now}')
        dr.font.size = Pt(9.5)
        dr.font.color.rgb = hex_rgb(GRAY)
        dr.italic = True

        doc.add_paragraph()

        # ── KPI Summary ──────────────────────────────
        max_disk = max(disk_data, key=lambda x: int(x.get('use_pct', 0))) \
            if disk_data else {'use_pct': 0, 'status': 'NORMAL'}
        add_kpi_row([
            ('CPU Usage', f"{cpu_data.get('usage', '--')}%",
             cpu_data.get('status', 'NORMAL')),
            ('Memory Usage', f"{mem_data.get('usage_pct', '--')}%",
             mem_data.get('status', 'NORMAL')),
            ('Disk Usage (peak)',
             f"{max_disk.get('use_pct', '--')}%", max_disk.get('status', 'NORMAL')),
        ])

        # ── 1. System Information ─────────────────────
        add_heading('1.  System Information')
        add_kv_table([
            ('Hostname', hostname),
            ('Operating System', os_name),
            ('Kernel Version', kernel),
            ('System Uptime', uptime),
            ('Report Date/Time', now),
        ])

        # ── 2. CPU ───────────────────────────────────
        add_heading('2.  CPU Utilization')
        add_kv_table([
            ('CPU Usage', f"{cpu_data.get('usage', '--')}%"),
            ('Load Average (1 min)', cpu_data.get('load_1min', '--')),
            ('Status', cpu_data.get('status', 'NORMAL')),
        ])

        # ── 3. Memory ────────────────────────────────
        add_heading('3.  Memory Usage')
        add_kv_table([
            ('Total RAM', f"{round(mem_data.get('total_mb', 0)/1024, 1)} GB"),
            ('Used RAM', f"{round(mem_data.get('used_mb', 0)/1024, 1)} GB"),
            ('Available', f"{round(mem_data.get('avail_mb', 0)/1024, 1)} GB"),
            ('Usage', f"{mem_data.get('usage_pct', '--')}%"),
            ('Status', mem_data.get('status', 'NORMAL')),
        ])

        # ── 4. Disk Table ────────────────────────────
        add_heading('4.  Disk / File System Usage')
        disk_headers = ['Filesystem', 'Size', 'Used',
                        'Available', 'Use%', 'Mount Point', 'Status']
        dtable = doc.add_table(rows=1, cols=7)
        make_header_row(dtable, disk_headers,
                        widths=[0.7, 0.5, 0.5, 0.6, 0.45, 2.75, 0.8])

        for fs in disk_data:
            row = dtable.add_row()
            st = fs.get('status', 'NORMAL')
            fg, bg = status_colors(st)
            vals = [fs.get('filesystem', ''), fs.get('size', ''), fs.get('used', ''),
                    fs.get('avail', ''), str(fs.get('use_pct', '')) + '%',
                    fs.get('mounted', ''), st]
            for i, val in enumerate(vals):
                cell = row.cells[i]
                is_status_col = (i == 6)
                align = WD_ALIGN_PARAGRAPH.CENTER if i in (
                    1, 2, 3, 4, 6) else WD_ALIGN_PARAGRAPH.LEFT
                cell_text(cell, val, size=9,
                          bold=is_status_col,
                          color=fg if is_status_col else '1A1A1A',
                          align=align)
                # Shade the whole row when not NORMAL, so problems jump out
                set_cell_bg(cell, bg if st != 'NORMAL' else 'FFFFFF')
        set_table_borders(dtable, color='D9D9D9', sz=4)

        # ── 5. Process Table ─────────────────────────
        add_heading('5.  Active Processes (Top 15 by CPU)')
        ptable = doc.add_table(rows=1, cols=4)
        make_header_row(ptable, ['PID', 'Process Name', 'CPU %', 'MEM %'],
                        widths=[0.8, 3.0, 1.0, 1.0])

        for idx, proc in enumerate(proc_data):
            row = ptable.add_row()
            bg = 'FFFFFF' if idx % 2 == 0 else ROW_ALT
            vals = [str(proc.get('pid', '')), proc.get('name', ''),
                    str(proc.get('cpu', '')) + '%', str(proc.get('mem', '')) + '%']
            for i, val in enumerate(vals):
                cell = row.cells[i]
                align = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
                cell_text(cell, val, size=9.5, color='1A1A1A', align=align)
                set_cell_bg(cell, bg)
        set_table_borders(ptable, color='D9D9D9', sz=4)

        # ── Footer ───────────────────────────────────
        doc.add_paragraph()
        fp = doc.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(10)
        fr = fp.add_run('── Generated by Linux System Monitoring Tool ──')
        fr.font.size = Pt(9)
        fr.font.color.rgb = hex_rgb(GRAY)
        fr.italic = True

        # ── Save & Send ──────────────────────────────
        timestamp = datetime.now().strftime('%Y-%m-%d_%H-%M')
        report_path = os.path.join(os.path.dirname(__file__), 'reports',
                                   f'system_report_{timestamp}.docx')
        os.makedirs(os.path.dirname(report_path), exist_ok=True)
        doc.save(report_path)
        return send_file(report_path, as_attachment=True,
                         download_name=f'system_report_{timestamp}.docx')

    except Exception as e:
        import traceback
        return jsonify({'error': str(e), 'trace': traceback.format_exc()}), 500

# Live Monitor Data (CPU + RAM + Disk combined)


@app.route('/api/live')
def live():
    try:
        import json
        cpu_data = json.loads(run_script('cpu.sh'))
        mem_data = json.loads(run_script('memory.sh'))
        disk_data = json.loads(run_script('disk.sh'))
        max_disk = max(disk_data, key=lambda x: int(x.get('use_pct', 0)))
        return jsonify({'cpu': cpu_data, 'memory': mem_data, 'disk': max_disk})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ===================================================
# RUN APP
# ===================================================
if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
